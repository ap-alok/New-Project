import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
rng = np.random.default_rng(42)
n = 1500
locs = {"City Center":11000,"Riverside":9500,"Tech Park":8800,"Old Town":7200,
        "Green Valley":6800,"Lake View":8200,"Industrial Area":4800,"Outskirts":3900}
loc = rng.choice(list(locs), n, p=[.10,.11,.16,.12,.14,.10,.12,.15])
bed = rng.choice([1,2,3,4,5], n, p=[.15,.38,.32,.12,.03])
area = np.clip(rng.normal(350 + bed*330, 120), 300, 4000).round()
bath = np.clip(bed - rng.choice([0,0,1], n), 1, None)
balc = rng.integers(0, 4, n)
total_fl = rng.choice([2,4,7,12,20,30], n)
floor = np.array([rng.integers(0, t+1) for t in total_fl])
age = rng.integers(0, 31, n)
furn = rng.choice(["Unfurnished","Semi-Furnished","Furnished"], n, p=[.35,.45,.20])
park = rng.choice([0,1,2], n, p=[.25,.55,.20])
metro = rng.choice(["Yes","No"], n, p=[.35,.65])
rate = np.array([locs[l] for l in loc])
price = (area*rate*(1 - 0.008*age)
         * np.select([furn=="Furnished", furn=="Semi-Furnished"], [1.10,1.04], 1.0)
         * np.where(metro=="Yes", 1.08, 1.0) * (1 + 0.03*park) * (1 + 0.004*floor)
         + bath*150000) * rng.lognormal(0, 0.10, n)
df = pd.DataFrame({"House_ID":[f"H{i:04d}" for i in range(1,n+1)],"Location":loc,
   "Area_sqft":area.astype(int),"Bedrooms":bed,"Bathrooms":bath,"Balconies":balc,
   "Floor":floor,"Total_Floors":total_fl,"Age_Years":age,"Furnishing":furn,
   "Parking_Spaces":park,"Near_Metro":metro,"Price_Lakhs":(price/1e5).round(2)})
# realistic messiness for cleaning practice
for col, k in [("Area_sqft",30),("Bathrooms",25),("Age_Years",40),("Furnishing",20)]:
    df.loc[rng.choice(n,k,replace=False), col] = np.nan
df = pd.concat([df, df.sample(12, random_state=1)], ignore_index=True)   # duplicates
df.loc[rng.choice(len(df),5,replace=False), "Price_Lakhs"] *= 8           # outliers
df.to_excel("data/house_prices.xlsx", sheet_name="House_Data", index=False)

dd = pd.DataFrame([
 ("House_ID","Unique property identifier","Text"),
 ("Location","Locality / zone of the property","Category (8 zones)"),
 ("Area_sqft","Carpet area in square feet","Number"),
 ("Bedrooms","Number of bedrooms (BHK)","Integer"),
 ("Bathrooms","Number of bathrooms","Integer"),
 ("Balconies","Number of balconies","Integer"),
 ("Floor","Floor the flat is on (0 = ground)","Integer"),
 ("Total_Floors","Total floors in the building","Integer"),
 ("Age_Years","Age of the property in years","Integer"),
 ("Furnishing","Furnishing status","Unfurnished / Semi-Furnished / Furnished"),
 ("Parking_Spaces","Number of parking spaces","Integer"),
 ("Near_Metro","Within walking distance of a metro station","Yes / No"),
 ("Price_Lakhs","TARGET - sale price in INR lakhs (1 lakh = 100,000)","Number"),
], columns=["Column","Description","Type"])
with pd.ExcelWriter("data/house_prices.xlsx", mode="a", engine="openpyxl", if_sheet_exists="overlay") as w:
    dd.to_excel(w, sheet_name="Data_Dictionary", index=False)
    pd.DataFrame({"Note":["Synthetic dataset generated for teaching/practice.",
      "Contains intentional missing values, 12 duplicate rows and a few price outliers for data-cleaning practice."]}
      ).to_excel(w, sheet_name="Data_Dictionary", index=False, startrow=16)
wb = load_workbook("data/house_prices.xlsx")
hdr = PatternFill("solid", fgColor="1F4E78")
for ws in wb:
    for row in ws.iter_rows():
        for c in row: c.font = Font(name="Arial", size=10)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = hdr
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col)+3, 70)
wb["Data_Dictionary"]["A17"].font = Font(name="Arial", bold=True)
wb.save("data/house_prices.xlsx")
print(df.shape); print(df.describe().T[["mean","min","max"]])
